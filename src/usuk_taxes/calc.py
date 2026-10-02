"""Generate the tax-year calculation workbook (Excel).

The workbook is a scaffold: blue cells are inputs, white cells are formulas,
grey cells are year constants (verify against the corpus before filing).
It models US TY 2025 and UK 2025-26 side by side, with the FTC limitation and
a FEIE-vs-FTC comparison. It is an estimation aid, not tax software.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

INPUT_FILL = PatternFill("solid", fgColor="DDEBF7")   # light blue: enter data
CONST_FILL = PatternFill("solid", fgColor="EDEDED")   # grey: year constants
HEAD_FILL = PatternFill("solid", fgColor="305496")    # dark blue header
HEAD_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
NOTE_FONT = Font(italic=True, color="666666")
THIN = Border(bottom=Side(style="thin", color="BBBBBB"))

USD = '"$"#,##0'
GBP = '"£"#,##0'
PCT = "0.0%"


def _head(ws, row, text, span=4):
    ws.cell(row=row, column=1, value=text).font = HEAD_FONT
    for col in range(1, span + 1):
        ws.cell(row=row, column=col).fill = HEAD_FILL
    return row + 1


def _label(ws, row, text, note=None):
    ws.cell(row=row, column=1, value=text)
    if note:
        ws.cell(row=row, column=3, value=note).font = NOTE_FONT


def _input(ws, row, label, value=None, fmt=GBP, note=None):
    _label(ws, row, label, note)
    c = ws.cell(row=row, column=2, value=value)
    c.fill = INPUT_FILL
    c.number_format = fmt
    return row + 1


def _const(ws, row, label, value, fmt=USD, note=None):
    _label(ws, row, label, note)
    c = ws.cell(row=row, column=2, value=value)
    c.fill = CONST_FILL
    c.number_format = fmt
    return row + 1


def _formula(ws, row, label, formula, fmt=USD, note=None, bold=False):
    _label(ws, row, label, note)
    c = ws.cell(row=row, column=2, value=formula)
    c.number_format = fmt
    if bold:
        c.font = BOLD
        ws.cell(row=row, column=1).font = BOLD
    return row + 1


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_workbook(path) -> None:
    wb = Workbook()

    # ── README ───────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "README"
    _widths(ws, [110])
    lines = [
        "US/UK expat tax workbook — US tax year 2025 / UK tax year 2025-26",
        "",
        "Cell convention:  BLUE = your inputs   GREY = year constants (verify!)   WHITE = calculated",
        "",
        "This is an estimation scaffold, not tax software. Grey constants were entered from 2025",
        "figures known at generation time — verify each against the corpus (`usuk ask ...`) or the",
        "official instructions before relying on a number.",
        "",
        "Sheets:",
        "  FX          — exchange rates (IRS yearly average for the US return; HMRC rates for SA)",
        "  Employment  — UK salary + PAYE from P60/payslips, converted to USD",
        "  Options     — per-exercise/vest rows with US/UK workday apportionment",
        "  Rental      — UK property: parallel UK (SA105) and US (Schedule E + ADS depreciation) computations",
        "  UK Tax      — rough UK 2025-26 liability (England bands; PAYE-adjusted)",
        "  US FTC      — Form 1116 general-basket limitation and excess-credit position",
        "  FEIE vs FTC — side-by-side strategy comparison",
        "",
        "Key deadlines:  US: 15 Jun 2026 (auto expat extension; 15 Oct with Form 4868)",
        "                UK: 31 Jan 2027 (online SA filing + balancing payment)",
        "",
        "Not professional advice — verify high-stakes positions with a cross-border professional.",
    ]
    for i, line in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=line)
    ws.cell(row=1, column=1).font = Font(bold=True, size=13)

    # ── FX ───────────────────────────────────────────────────────────────
    ws = wb.create_sheet("FX")
    _widths(ws, [46, 14, 60])
    r = _head(ws, 1, "Exchange rates", 3)
    r = _const(ws, r, "IRS yearly average 2025 (GBP per 1 USD)", 0.759, "0.000",
               "irs.gov 'Yearly average currency exchange rates' — use for US return income")
    r = _const(ws, r, "IRS yearly average 2024 (GBP per 1 USD)", 0.783, "0.000", "for reference / prior-year items")
    r = _input(ws, r, "Spot rate override (GBP per 1 USD), optional", None, "0.000",
               "specific-transaction rate (e.g. option exercise date) if you prefer spot for that item")
    ws.cell(row=r, column=1, value="USD = GBP / rate.  UK SA is computed in GBP directly (no conversion needed).").font = NOTE_FONT
    # Named-style anchor: other sheets reference FX!B2
    fx = "FX!$B$2"

    # ── Employment ───────────────────────────────────────────────────────
    ws = wb.create_sheet("Employment")
    _widths(ws, [46, 16, 66])
    r = _head(ws, 1, "UK employment income (P60 / final 2025 payslips)", 3)
    ws.cell(row=r, column=3, value="US return needs CALENDAR-2025 amounts; UK SA needs 2025-26 amounts. "
            "Enter both — payslips are the bridge.").font = NOTE_FONT
    r += 1
    r = _input(ws, r, "Gross pay, calendar 2025 (GBP)", None, GBP, "sum Jan-Dec 2025 payslips")
    r = _input(ws, r, "UK tax (PAYE) withheld, calendar 2025 (GBP)", None, GBP, "accrued/paid basket choice on Form 1116 — be consistent")
    r = _input(ws, r, "Gross pay, UK year 2025-26 (GBP)", None, GBP, "P60 for 2025-26")
    r = _input(ws, r, "UK tax (PAYE) withheld, 2025-26 (GBP)", None, GBP)
    r = _input(ws, r, "Employee NIC, 2025-26 (GBP)", None, GBP, "NIC is NOT creditable as income tax on Form 1116")
    r += 1
    r = _formula(ws, r, "Gross pay, calendar 2025 (USD)", f"=IF(B3>0,B3/{fx},0)", USD, bold=True)
    r = _formula(ws, r, "UK tax withheld, calendar 2025 (USD)", f"=IF(B4>0,B4/{fx},0)", USD, bold=True)

    # ── Options ──────────────────────────────────────────────────────────
    ws = wb.create_sheet("Options")
    _widths(ws, [18, 12, 12, 12, 10, 11, 11, 13, 13, 13, 14, 14, 14])
    ws.cell(row=1, column=1, value="Equity compensation — one row per taxable event (NQSO exercise / RSU vest / ISO exercise)").font = BOLD
    ws.cell(row=2, column=1, value="Sourcing: spread is apportioned by workdays over grant→vest. UK workday share taxed by UK "
            "(usually via PAYE, see EIM43501/ERSM162000); total spread is US-taxable (citizen) with FTC for the UK share. "
            "ISO exercises: no regular-tax spread but AMT adjustment (Form 6251).").font = NOTE_FONT
    headers = ["Grant/award ID", "Type", "Grant date", "Vest/exercise date", "Shares", "Strike $", "FMV $",
               "Spread $ (calc)", "Days grant→vest", "UK workdays", "UK share (calc)", "UK-source $ (calc)", "US-source $ (calc)"]
    hr = 4
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=hr, column=col, value=h)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in range(hr + 1, hr + 11):
        for col in (1, 2, 3, 4, 5, 6, 7, 9, 10):
            ws.cell(row=row, column=col).fill = INPUT_FILL
        ws.cell(row=row, column=8, value=f"=IF(E{row}>0,E{row}*(G{row}-F{row}),0)").number_format = USD
        ws.cell(row=row, column=11, value=f"=IF(I{row}>0,J{row}/I{row},0)").number_format = PCT
        ws.cell(row=row, column=12, value=f"=H{row}*K{row}").number_format = USD
        ws.cell(row=row, column=13, value=f"=H{row}-L{row}").number_format = USD
    tr = hr + 11
    ws.cell(row=tr, column=7, value="Totals:").font = BOLD
    for col, letter in ((8, "H"), (12, "L"), (13, "M")):
        c = ws.cell(row=tr, column=col, value=f"=SUM({letter}{hr+1}:{letter}{tr-1})")
        c.font = BOLD
        c.number_format = USD
    ws.cell(row=tr + 2, column=1, value="UK tax paid via PAYE on option/RSU income, calendar 2025 (GBP):")
    ws.cell(row=tr + 2, column=3).fill = INPUT_FILL
    ws.cell(row=tr + 2, column=3).number_format = GBP

    # ── Rental ───────────────────────────────────────────────────────────
    ws = wb.create_sheet("Rental")
    _widths(ws, [46, 16, 16, 56])
    ws.cell(row=1, column=1, value="UK rental property — parallel computations").font = BOLD
    ws.cell(row=2, column=1, value="UK column: 2025-26 amounts, SA105 rules (no depreciation; finance-cost 20% credit). "
            "US column: calendar-2025 amounts in GBP (converted below), Schedule E rules (depreciation; full mortgage interest).").font = NOTE_FONT
    ws.cell(row=4, column=2, value="UK 2025-26 (GBP)").font = BOLD
    ws.cell(row=4, column=3, value="US cal-2025 (GBP)").font = BOLD
    rows = [
        ("Rents received", None),
        ("Letting agent fees", None),
        ("Repairs and maintenance (not improvements)", None),
        ("Insurance", None),
        ("Other allowable expenses", None),
        ("Mortgage interest (finance costs)", "UK: NOT deducted here — 20% credit below. US: fully deductible."),
    ]
    r = 5
    for label, note in rows:
        ws.cell(row=r, column=1, value=label)
        for col in (2, 3):
            ws.cell(row=r, column=col).fill = INPUT_FILL
            ws.cell(row=r, column=col).number_format = GBP
        if note:
            ws.cell(row=r, column=4, value=note).font = NOTE_FONT
        r += 1
    int_row = r - 1
    r = _formula(ws, r, "UK property profit (before finance costs)", f"=B5-SUM(B6:B9)", GBP, "→ SA105", True)
    ws.cell(row=r - 1, column=1, value="UK property profit (excl. finance costs)")
    r = _formula(ws, r, "UK finance-cost tax credit (20%)", f"=0.2*B{int_row}", GBP, "reduces UK tax, not profit")
    r += 1
    r = _input(ws, r, "US depreciable basis (building only, USD)", None, USD,
               "lower of cost or FMV when first rented, excl. land; foreign residential = 30-yr ADS straight line")
    basis_row = r - 1
    r = _formula(ws, r, "US annual depreciation (30-yr ADS)", f"=B{basis_row}/30", USD)
    dep_row = r - 1
    r = _formula(ws, r, "US Schedule E net income (USD)",
                 f"=(C5-SUM(C6:C10))/{fx}-B{dep_row}", USD, "rents less ALL expenses incl. interest, less depreciation", True)

    # ── UK Tax ───────────────────────────────────────────────────────────
    ws = wb.create_sheet("UK Tax")
    _widths(ws, [46, 16, 66])
    r = _head(ws, 1, "UK 2025-26 liability estimate (England/NI bands)", 3)
    r = _const(ws, r, "Personal allowance", 12570, GBP, "tapered £1 per £2 of income over £100,000")
    r = _const(ws, r, "Basic-rate band (20%)", 37700, GBP)
    r = _const(ws, r, "Additional-rate threshold (45%)", 125140, GBP)
    r += 1
    r = _formula(ws, r, "Employment income 2025-26", "=Employment!B5", GBP)
    r = _formula(ws, r, "Property profit", "='Rental'!B11", GBP)
    r = _input(ws, r, "Other UK-taxable income (interest, dividends...)", None, GBP, "model dividend rates separately if material")
    ti_first, ti_last = r - 3, r - 1
    r = _formula(ws, r, "Total income", f"=SUM(B{ti_first}:B{ti_last})", GBP, bold=True)
    tot = r - 1
    r = _formula(ws, r, "Personal allowance (after taper)",
                 f"=MAX(0,B2-MAX(0,(B{tot}-100000)/2))", GBP)
    pa = r - 1
    r = _formula(ws, r, "Taxable income", f"=MAX(0,B{tot}-B{pa})", GBP, bold=True)
    tax_base = r - 1
    r = _formula(ws, r, "Tax at 20% / 40% / 45%",
                 f"=0.2*MIN(B{tax_base},B3)+0.4*MAX(0,MIN(B{tax_base},B4)-B3)+0.45*MAX(0,B{tax_base}-B4)",
                 GBP, "bands apply to taxable income; approximation — bands shift with pension/gift-aid contributions")
    gross_tax = r - 1
    r = _formula(ws, r, "Less: finance-cost credit (rental)", "='Rental'!B12", GBP)
    r = _formula(ws, r, "Estimated UK liability", f"=MAX(0,B{gross_tax}-B{gross_tax+1})", GBP, bold=True)
    liab = r - 1
    r = _formula(ws, r, "Less: PAYE paid 2025-26", "=Employment!B6", GBP)
    r = _formula(ws, r, "Balancing payment due 31 Jan 2027 (est.)", f"=B{liab}-B{liab+1}", GBP, bold=True)

    # ── US FTC ───────────────────────────────────────────────────────────
    ws = wb.create_sheet("US FTC")
    _widths(ws, [52, 16, 66])
    r = _head(ws, 1, "Form 1116 — general category limitation (calendar 2025, USD)", 3)
    r = _const(ws, r, "Standard deduction (MFJ, 2025)", 31500, USD, "verify in i1040; itemize instead if better")
    r += 1
    r = _formula(ws, r, "Foreign-source wages", "=Employment!B9", USD)
    r = _formula(ws, r, "Foreign-source option/RSU income (UK share)", "=Options!L15", USD,
                 "US-workday share is US-source and NOT in the FTC numerator")
    r = _formula(ws, r, "Foreign-source rental net income", "='Rental'!B16", USD, "passive basket if net positive — see note")
    fs_first, fs_last = r - 3, r - 1
    r = _formula(ws, r, "Foreign-source taxable income (FSTI)", f"=SUM(B{fs_first}:B{fs_last})", USD, bold=True)
    fsti = r - 1
    r = _input(ws, r, "US-source income (US option share, US interest...)", None, USD)
    r = _formula(ws, r, "Total taxable income before deduction", f"=B{fsti}+B{fsti+1}", USD)
    r = _formula(ws, r, "Taxable income (after standard deduction)", f"=MAX(0,B{fsti+2}-B2)", USD, bold=True)
    ti = r - 1
    r = _input(ws, r, "US tax before credits (from 1040 tables on line above)", None, USD,
               "compute from 2025 MFJ brackets; then subtract CTC/ODC AFTER the FTC per form ordering")
    ustax = r - 1
    r = _formula(ws, r, "FTC limitation = US tax × FSTI / TI", f"=IF(B{ti}>0,B{ustax}*MIN(1,B{fsti}/B{ti}),0)", USD, bold=True)
    lim = r - 1
    r = _formula(ws, r, "UK tax paid/accrued on that income (USD)",
                 f"=Employment!B10+IF(Options!C17>0,Options!C17/{fx},0)", USD,
                 "employment PAYE + option PAYE; add UK tax on rental when SA is finalised")
    paid = r - 1
    r = _formula(ws, r, "Credit usable this year", f"=MIN(B{lim},B{paid})", USD, bold=True)
    r = _formula(ws, r, "Excess (carryback 1 yr / carryforward 10 yrs)", f"=MAX(0,B{paid}-B{lim})", USD)
    ws.cell(row=r + 1, column=1, value="Note: net rental income is usually PASSIVE basket (separate 1116). UK tax rates exceed US "
            "rates, so high-taxed passive income may kick back to general — model when figures are real.").font = NOTE_FONT

    # ── FEIE vs FTC ──────────────────────────────────────────────────────
    ws = wb.create_sheet("FEIE vs FTC")
    _widths(ws, [52, 18, 18, 52])
    r = _head(ws, 1, "Strategy comparison — decide with real numbers", 4)
    r = _const(ws, r, "FEIE limit per spouse (2025)", 130000, USD, "verify in i2555")
    ws.cell(row=4, column=2, value="A: FTC only").font = BOLD
    ws.cell(row=4, column=3, value="B: FEIE + FTC").font = BOLD
    considerations = [
        ("Earned income covered", "credit for UK tax", "excluded up to limit"),
        ("Child Tax Credit (refundable ACTC)", "available", "LOST if income fully excluded — often decisive with 2-3 qualifying kids"),
        ("Unexcluded income tax rate", "n/a", "stacking rule: taxed at higher brackets"),
        ("Carryovers build up", "yes (10-yr carryforward)", "no carryover on excluded income"),
        ("Revocation trap", "none", "revoking FEIE bars re-election for 5 years without IRS consent"),
        ("Typical outcome, high UK tax + kids + rental", "usually better", "rarely better"),
    ]
    r = 5
    for label, a, b in considerations:
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=a)
        ws.cell(row=r, column=3, value=b)
        for col in range(1, 4):
            ws.cell(row=r, column=col).border = THIN
            ws.cell(row=r, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    ws.cell(row=r + 1, column=1, value="UK marginal rates (20/40/45%) generally exceed US rates, so the FTC usually wipes out US tax "
            "on UK earnings AND preserves the refundable Child Tax Credit. Verify by filling both scenarios once real numbers exist.").font = NOTE_FONT

    wb.save(path)
